from __future__ import annotations

import importlib
import json
import sqlite3
import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import ontology_platform.automation as automation_module
from ontology_platform.adapters import get_adapter
from ontology_platform.automation import execute_operation, preflight_operation
from ontology_platform.contract_documents import parse_rule_docx_bytes
from ontology_platform.coverage import build_semantic_coverage
from ontology_platform.database import connect, initialize_platform_db
from ontology_platform.decisions import list_decisions
from ontology_platform.events import EventType, declare_event_type, init_event_schema, record_event
from ontology_platform.governance import (
    bulk_review_semantic_mappings,
    derive_ontology_version,
    list_business_rules,
    list_semantic_mappings,
    publish_ontology,
    review_semantic_mapping,
    upsert_business_rule,
)
from ontology_platform.industry_blueprints import (
    infer_industry_blueprint,
    list_industry_blueprints,
    upsert_industry_blueprint,
)
from ontology_platform.kernel_package import build_kernel_package, export_kernel_package
from ontology_platform.metadata import (
    analyze_schema_drift,
    assess_data_source_readiness,
    check_business_api_gateway,
    check_data_source_connection,
    import_openapi_operations,
    import_openapi_operations_from_url,
    list_data_sources,
    list_source_apis,
    register_data_source,
    register_source_api,
    scan_data_source,
)
from ontology_platform.model_client import (
    OpenRouterClient,
    OpenRouterConfig,
    generate_blueprint_draft,
    generate_ontology_reasoning_chain,
    get_model_config,
    reset_model_config,
    update_model_config,
)
from ontology_platform.natural_language import query_natural_language
from ontology_platform.onboarding import run_onboarding_pipeline
from ontology_platform.ontology import explain_instance, export_ontology_asset, generate_ontology_draft, list_ontologies
from ontology_platform.operation_bindings import assess_operation_bindings
from ontology_platform.release_readiness import assess_ontology_release_readiness
from ontology_platform.sample_data import create_contract_sample_db, create_equipment_sample_db
from ontology_platform.semantic_kernel import assess_decision_consistency, assess_instance, list_instance_ids
from ontology_platform.temporal import init_temporal_schema, record_attribute_version


def test_metadata_to_ontology_flow(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan = scan_data_source(platform_db, source.id)

    assert len(scan["tables"]) == 4

    ontology = generate_ontology_draft(platform_db, source.id)
    object_codes = {item["code"] for item in ontology["objects"]}
    assert {"contract", "customer", "payment_plan", "invoice"}.issubset(object_codes)
    assert len(ontology["relations"]) >= 3
    assert any(rule["code"] == "blacklist_customer_warning" for rule in ontology["rules"])

    explanation = explain_instance(platform_db, ontology["id"], "contract", "1")
    assert explanation["object"]["name"] == "合同"
    values = {item["attributeCode"]: item["value"] for item in explanation["attributes"]}
    assert values["contract_no"] == "HT-2026-001"
    assert values["amount"] == 1200000

    assessment = assess_instance(platform_db, ontology["id"], "contract", "3")
    assert assessment["decision"]["status"] == "review"
    failed_rules = {item["ruleCode"] for item in assessment["ruleResults"] if not item["passed"]}
    assert {"blacklist_customer_warning", "payment_plan_amount_match"}.issubset(failed_rules)

    preflight = preflight_operation(platform_db, ontology["id"], source.id, "submit_contract", "3")
    assert preflight["allowed"] is False
    assert preflight["nextAction"] == "route_to_human_review"


def test_natural_language_query_routes_to_semantic_kernel(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    compliance = query_natural_language(platform_db, "合同 1 是否合规？", ontology["id"], source.id, use_model=False)
    assert compliance["intent"] == "compliance_assessment"
    assert compliance["resolved"]["objectCode"] == "contract"
    assert compliance["resolved"]["instanceId"] == "1"
    assert compliance["evidence"]["decision"]["status"] == "approved"
    assert "合规" in compliance["answer"]

    overall = query_natural_language(platform_db, "合同是否合规？", ontology["id"], source.id, use_model=False)
    assert overall["intent"] == "compliance_assessment"
    assert overall["evidence"]["objectCode"] == "contract"
    assert overall["evidence"]["assessed"] > 0
    assert "已评估" in overall["answer"]

    preflight = query_natural_language(
        platform_db, "HT-2026-003 能提交审批吗？", ontology["id"], source.id, use_model=False
    )
    assert preflight["intent"] == "operation_preflight"
    assert preflight["resolved"]["instanceId"] == "3"
    assert preflight["resolved"]["operationCode"] == "submit_contract"
    assert preflight["evidence"]["allowed"] is False
    assert "不建议自动执行" in preflight["answer"]


def test_rule_word_document_can_define_custom_rules(tmp_path: Path) -> None:
    from io import BytesIO

    from docx import Document

    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    doc = Document()
    doc.add_paragraph("合同自定义规则")
    table = doc.add_table(rows=1, cols=7)
    headers = ["规则编码", "规则名称", "适用对象", "规则类型", "规则表达式", "严重程度", "自然语言说明"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    row = table.add_row().cells
    values = [
        "contract_title_required_word",
        "合同标题必填（Word导入）",
        "contract",
        "校验",
        "title != null",
        "阻断",
        "合同标题必须填写，不能为空。",
    ]
    for index, value in enumerate(values):
        row[index].text = value
    buffer = BytesIO()
    doc.save(buffer)

    parsed = parse_rule_docx_bytes("rules.docx", buffer.getvalue())
    assert len(parsed["rules"]) == 1
    assert parsed["rules"][0]["code"] == "contract_title_required_word"
    assert parsed["rules"][0]["ruleType"] == "validation"
    assert parsed["rules"][0]["severity"] == "blocking"

    rule = parsed["rules"][0]
    upsert_business_rule(
        platform_db,
        ontology["id"],
        rule["code"],
        rule["name"],
        rule["ruleType"],
        rule["scopeObjectCode"],
        rule["expression"],
        rule["severity"],
        rule["naturalLanguage"],
        "test",
        rule["status"],
    )
    assert any(
        rule["code"] == "contract_title_required_word"
        for rule in list_business_rules(platform_db, ontology["id"])["items"]
    )


def test_operation_execution_respects_semantic_preflight(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    ready = execute_operation(
        platform_db,
        ontology["id"],
        source.id,
        "submit_contract",
        "1",
        payload={"comment": "自动提交"},
        actor="测试用户",
    )
    assert ready["status"] == "ready_for_execution"
    assert ready["executed"] is False
    assert ready["preflight"]["allowed"] is True
    assert ready["preflight"]["decision"]["decisionId"].startswith("DR-")
    assert ready["decisionRecord"]["decisionId"].startswith("DR-")
    assert ready["execution"]["path"] == "/contracts/1/submit"
    assert ready["execution"]["payload"] == {"comment": "自动提交"}

    blocked = execute_operation(platform_db, ontology["id"], source.id, "submit_contract", "3", actor="测试用户")
    assert blocked["status"] == "blocked_by_semantic_kernel"
    assert blocked["executed"] is False
    assert blocked["preflight"]["allowed"] is False
    assert blocked["decisionRecord"]["decisionType"] == "operation_execution"

    with connect(platform_db) as conn:
        audit_count = conn.execute("select count(*) from audit_log where action = 'execute_operation'").fetchone()[0]
    assert audit_count == 2

    decisions = list_decisions(platform_db, 10)
    decision_types = [item["decisionType"] for item in decisions]
    assert decision_types.count("instance_assessment") == 2
    assert decision_types.count("operation_preflight") == 2
    assert decision_types.count("operation_execution") == 2


def test_real_operation_execution_uses_business_api_base_url(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db,
        "合同管理样例系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        api_base_url="https://legacy.example/api",
        api_headers={"Authorization": "Bearer legacy-token", "X-Tenant": "contract"},
    )
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")
    calls: list[dict[str, object]] = []

    def fake_invoke(
        base_url: str, execution_plan: dict[str, object], timeout_seconds: float, headers: dict[str, str] | None = None
    ) -> dict[str, object]:
        calls.append({"baseUrl": base_url, "plan": execution_plan, "timeout": timeout_seconds, "headers": headers})
        return {"statusCode": 200, "body": {"accepted": True}}

    original_invoke = automation_module._invoke_http_operation
    automation_module._invoke_http_operation = fake_invoke
    try:
        result = execute_operation(
            platform_db,
            ontology["id"],
            source.id,
            "submit_contract",
            "1",
            payload={"comment": "自动提交"},
            actor="测试用户",
            dry_run=False,
            timeout_seconds=7,
        )
    finally:
        automation_module._invoke_http_operation = original_invoke

    assert result["status"] == "executed"
    assert result["executed"] is True
    assert calls[0]["baseUrl"] == "https://legacy.example/api"
    assert calls[0]["timeout"] == 7
    assert calls[0]["headers"] == {"Authorization": "Bearer legacy-token", "X-Tenant": "contract"}
    assert calls[0]["plan"]["path"] == "/contracts/1/submit"
    assert result["execution"]["remote"]["body"] == {"accepted": True}


def test_decision_consistency_batch_assesses_samples_and_reports_rule_distribution(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    instance_ids = list_instance_ids(platform_db, ontology["id"], "contract", 2)
    assert [str(item) for item in instance_ids] == ["1", "2"]

    report = assess_decision_consistency(platform_db, ontology["id"], "contract", limit=3)

    assert report["objectCode"] == "contract"
    assert report["sampleSize"] == 3
    assert report["assessed"] == 3
    assert report["summary"]["approved"] == 2
    assert report["summary"]["review"] == 1
    assert report["status"] == "mixed"
    assert {item["instanceId"] for item in report["items"]} == {"1", "2", "3"}
    assert report["ruleFailures"][0]["ruleCode"] in {"blacklist_customer_warning", "payment_plan_amount_match"}
    assert report["nextActions"]

    explicit = assess_decision_consistency(platform_db, ontology["id"], "contract", ["1", "missing"], 10)
    assert explicit["assessed"] == 1
    assert explicit["errorCount"] == 1
    assert explicit["status"] == "incomplete"

    with connect(platform_db) as conn:
        audit_count = conn.execute(
            "select count(*) from audit_log where action = 'assess_decision_consistency'"
        ).fetchone()[0]
    assert audit_count == 2


def test_data_source_readiness_identifies_onboarding_gaps(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")

    initial = assess_data_source_readiness(platform_db, source.id)
    assert initial["status"] == "blocked"
    assert initial["summary"]["tables"] == 0
    assert {"metadata_scan", "ontology", "business_apis"}.issubset({item["code"] for item in initial["gaps"]})

    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")
    reviewed = assess_data_source_readiness(platform_db, source.id)

    assert reviewed["score"] >= 65
    assert reviewed["status"] == "partial"
    assert reviewed["summary"]["tables"] == 4
    assert reviewed["summary"]["apis"] == 1
    assert reviewed["summary"]["ontologies"] == 1
    assert reviewed["summary"]["pendingMappings"] > 0
    assert any(item["code"] == "mapping_governance" for item in reviewed["gaps"])

    bulk_review_semantic_mappings(platform_db, ontology["id"], "confirmed", "业务专家", "接入准备度验证")
    ready = assess_data_source_readiness(platform_db, source.id)
    assert ready["score"] == 100
    assert ready["status"] == "ready"
    assert ready["gaps"] == []


def test_database_api_readiness_requires_business_api_base_url(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db,
        "合同管理 API 系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
    )
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")
    bulk_review_semantic_mappings(platform_db, ontology["id"], "confirmed", "业务专家", "API 基址门禁验证")

    missing_gateway = assess_data_source_readiness(platform_db, source.id)
    assert missing_gateway["status"] == "partial"
    assert missing_gateway["summary"]["requiresApiGateway"] is True
    assert missing_gateway["summary"]["apiBaseUrlConfigured"] is False
    assert any(item["code"] == "api_gateway" for item in missing_gateway["gaps"])

    updated = register_data_source(
        platform_db,
        "合同管理 API 系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        api_base_url="https://legacy.example/api",
    )
    ready = assess_data_source_readiness(platform_db, updated.id)

    assert ready["score"] == 100
    assert ready["status"] == "ready"
    assert ready["summary"]["apiBaseUrlConfigured"] is True
    assert ready["gaps"] == []


def test_onboarding_pipeline_registers_scans_generates_ontology_and_reports_readiness(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)

    result = run_onboarding_pipeline(
        platform_db,
        "合同管理一键接入",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        blueprint_id="contract-management",
    )

    assert result["connection"]["reachable"] is True
    assert result["apiGateway"]["status"] == "not_configured"
    assert result["scan"]["tables"]
    assert result["ontology"]["blueprint"]["id"] == "contract-management"
    assert result["readiness"]["summary"]["tables"] == 4
    assert result["status"] == "partial"
    assert [step["code"] for step in result["steps"]] == [
        "register_data_source",
        "test_connection",
        "test_api_gateway",
        "scan_metadata",
        "import_openapi",
        "generate_ontology",
        "seed_permissions",
        "assess_readiness",
    ]


def test_onboarding_pipeline_imports_openapi_operations(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    result = run_onboarding_pipeline(
        platform_db,
        "合同 API 一键接入",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        blueprint_id="contract-management",
        openapi_spec={
            "openapi": "3.0.0",
            "paths": {
                "/contracts/{id}/submit": {
                    "post": {
                        "operationId": "submit_contract",
                        "summary": "提交合同审批",
                        "x-semantic-action": "contract.submit_for_approval",
                        "responses": {"200": {"description": "ok"}},
                    }
                }
            },
        },
    )

    assert result["apiImport"]["count"] == 1
    assert any(step["code"] == "import_openapi" and step["status"] == "completed" for step in result["steps"])
    assert list_source_apis(platform_db, result["dataSource"]["id"])[0]["operation_code"] == "submit_contract"
    assert result["readiness"]["summary"]["apis"] == 1


def test_onboarding_pipeline_stops_when_connection_fails(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    result = run_onboarding_pipeline(
        platform_db,
        "不存在的传统系统",
        "sqlite",
        str(tmp_path / "missing.sqlite3"),
        domain="通用业务",
    )

    assert result["status"] == "blocked"
    assert result["connection"]["reachable"] is False
    assert result["scan"] is None
    assert result["ontology"] is None
    assert [step["code"] for step in result["steps"]] == ["register_data_source", "test_connection"]


def test_openapi_import_registers_business_operations(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db, "合同 API 系统", "sqlite", str(legacy_db), domain="合同管理", system_category="database+api"
    )
    spec = {
        "openapi": "3.0.0",
        "paths": {
            "/contracts/{id}/submit": {
                "post": {
                    "operationId": "submit_contract",
                    "summary": "提交合同审批",
                    "x-semantic-action": "contract.submit_for_approval",
                    "requestBody": {
                        "content": {
                            "application/json": {
                                "schema": {"type": "object", "properties": {"comment": {"type": "string"}}}
                            }
                        }
                    },
                    "responses": {"200": {"content": {"application/json": {"schema": {"type": "object"}}}}},
                }
            },
            "/contracts/{id}/archive": {
                "post": {
                    "summary": "归档合同",
                    "responses": {"202": {"description": "accepted"}},
                }
            },
        },
    }

    result = import_openapi_operations(platform_db, source.id, spec)
    apis = list_source_apis(platform_db, source.id)

    assert result["count"] == 2
    assert {api["operation_code"] for api in apis} == {"submit_contract", "post_contracts_id_archive"}
    submit = next(api for api in apis if api["operation_code"] == "submit_contract")
    archive = next(api for api in apis if api["operation_code"] == "post_contracts_id_archive")
    assert submit["semantic_action"] == "contract.submit_for_approval"
    assert json.loads(submit["request_schema"])["properties"]["comment"]["type"] == "string"
    assert archive["semantic_action"].startswith("archive.")

    readiness = assess_data_source_readiness(platform_db, source.id)
    assert readiness["summary"]["apis"] == 2


def test_openapi_import_from_url_uses_business_api_headers(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"
    seen_headers: dict[str, str] = {}
    spec = {
        "openapi": "3.0.0",
        "paths": {
            "/contracts/{id}/submit": {
                "post": {
                    "operationId": "submit_contract",
                    "summary": "提交合同审批",
                    "x-semantic-action": "contract.submit_for_approval",
                    "responses": {"200": {"description": "ok"}},
                }
            }
        },
    }

    class OpenApiHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            seen_headers["Authorization"] = self.headers.get("Authorization", "")
            seen_headers["X-Tenant"] = self.headers.get("X-Tenant", "")
            body = json.dumps(spec).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, format: str, *args: object) -> None:
            return None

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db,
        "合同 API 系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        api_headers={"Authorization": "Bearer openapi-token", "X-Tenant": "contract"},
    )
    server = HTTPServer(("127.0.0.1", 0), OpenApiHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/openapi.json"
        result = import_openapi_operations_from_url(platform_db, source.id, url)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)

    assert result["sourceUrl"] == url
    assert result["count"] == 1
    assert seen_headers["Authorization"] == "Bearer openapi-token"
    assert seen_headers["X-Tenant"] == "contract"
    assert list_source_apis(platform_db, source.id)[0]["operation_code"] == "submit_contract"


def test_kernel_package_exports_installable_semantic_kernel_manifest(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db,
        "合同管理语义内核系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        api_headers={"Authorization": "Bearer legacy-token"},
    )
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    package = build_kernel_package(platform_db, source.id, "https://kernel.example/api")
    asset = export_kernel_package(platform_db, source.id, "https://kernel.example/api")

    assert package["packageType"] == "ontology-semantic-kernel"
    assert package["dataSource"]["id"] == source.id
    assert package["dataSource"]["apiHeadersConfigured"] is True
    assert package["dataSource"]["apiHeaderNames"] == ["Authorization"]
    assert package["readiness"]["summary"]["tables"] == 4
    assert package["ontologies"][0]["id"] == ontology["id"]
    assert any(item["code"] == "contract" for item in package["ontologies"][0]["objects"])
    assert any(rule["code"] == "blacklist_customer_warning" for rule in package["ontologies"][0]["rules"])
    assert package["operations"][0]["operationCode"] == "submit_contract"
    assert package["operations"][0]["requiresPreflight"] is True
    assert (
        package["runtimeEndpoints"]["execute"]
        == "https://kernel.example/api/automation/operations/{operationCode}/execute"
    )
    assert package["governanceGates"]["executionRequiresApprovedDecision"] is True
    assert asset["filename"] == f"semantic-kernel-datasource-{source.id}.json"
    assert json.loads(asset["content"])["packageType"] == "ontology-semantic-kernel"


def test_semantic_coverage_reports_object_rule_mapping_and_operation_readiness(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")

    not_modeled = build_semantic_coverage(platform_db, source.id)
    assert not_modeled["status"] == "not_modeled"
    assert not_modeled["score"] == 0

    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    pending = build_semantic_coverage(platform_db, source.id)
    contract_pending = next(item for item in pending["objects"] if item["objectCode"] == "contract")
    assert pending["status"] in {"partial", "blocked"}
    assert contract_pending["pendingMappings"] > 0
    assert contract_pending["ruleCount"] > 0
    assert contract_pending["operationCount"] == 1
    assert contract_pending["automationReady"] is False

    bulk_review_semantic_mappings(platform_db, ontology["id"], "confirmed", "业务专家", "语义覆盖度验证")

    ready = build_semantic_coverage(platform_db, source.id)
    contract_ready = next(item for item in ready["objects"] if item["objectCode"] == "contract")
    assert ready["score"] > pending["score"]
    assert ready["summary"]["confirmedMappings"] > 0
    assert ready["summary"]["semanticOperations"] == 1
    assert ready["summary"]["fullyCoveredObjects"] >= 1
    assert contract_ready["automationReady"] is True
    assert any(item["operationCode"] == "submit_contract" for item in contract_ready["operations"])


def test_operation_binding_assessment_validates_semantic_actions_against_ontology_objects(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    register_source_api(
        platform_db,
        source.id,
        "submit_contract",
        "提交合同审批",
        "POST",
        "/contracts/{id}/submit",
        "contract.submit_for_approval",
    )
    register_source_api(
        platform_db, source.id, "archive_case", "归档案件", "POST", "/cases/{id}/archive", "case.archive"
    )
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")
    bulk_review_semantic_mappings(platform_db, ontology["id"], "confirmed", "业务专家", "API 绑定验证")

    bindings = assess_operation_bindings(platform_db, source.id)
    submit = next(item for item in bindings["items"] if item["operationCode"] == "submit_contract")
    archive = next(item for item in bindings["items"] if item["operationCode"] == "archive_case")

    assert bindings["status"] == "partial"
    assert bindings["summary"]["operations"] == 2
    assert bindings["summary"]["readyOperations"] == 1
    assert bindings["summary"]["unboundOperations"] == 1
    assert submit["status"] == "ready"
    assert submit["automationReady"] is True
    assert archive["status"] == "unbound"
    assert "未绑定到当前数据源本体对象" in archive["gaps"][0]


def test_industry_blueprints_seed_ontology_names_mappings_and_rules(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_equipment.sqlite3"

    blueprints = list_industry_blueprints()
    assert {item["id"] for item in blueprints}.issuperset(
        {"contract-management", "equipment-maintenance", "generic-enterprise"}
    )
    assert infer_industry_blueprint(["equipment", "work_order"]).id == "equipment-maintenance"

    initialize_platform_db(platform_db)
    create_equipment_sample_db(legacy_db)
    source = register_data_source(platform_db, "设备运维生产系统", "sqlite", str(legacy_db), domain="设备运维")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="equipment-maintenance")

    assert ontology["blueprint"]["id"] == "equipment-maintenance"
    assert any(item["code"] == "equipment" and item["name"] == "设备" for item in ontology["objects"])
    assert any(item["code"] == "criticality" and item["name"] == "重要等级" for item in ontology["attributes"])
    assert any(rule["code"] == "critical_equipment_open_fault" for rule in ontology["rules"])
    assert any(
        "设备运维蓝图" in mapping["evidence"] and mapping["confidence"] >= 0.92 for mapping in ontology["mappings"]
    )

    try:
        generate_ontology_draft(platform_db, source.id, name="不存在蓝图本体", blueprint_id="missing-blueprint")
    except ValueError as error:
        assert "行业蓝图不存在" in str(error)
    else:
        raise AssertionError("不存在的行业蓝图应被拒绝")


def test_custom_industry_blueprint_can_be_imported_and_used_for_generation(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_equipment.sqlite3"

    initialize_platform_db(platform_db)
    create_equipment_sample_db(legacy_db)
    blueprint = upsert_industry_blueprint(
        platform_db,
        {
            "id": "lab-asset",
            "name": "实验室资产蓝图",
            "domain": "实验室资产",
            "description": "覆盖实验室设备资产管理。",
            "objectHints": {"equipment": "实验设备"},
            "attributeHints": {"equipment_code": "资产编号", "criticality": "资产等级"},
            "rules": [
                {
                    "code": "lab_asset_code_required",
                    "name": "资产编号必填",
                    "rule_type": "validation",
                    "scope_object_code": "equipment",
                    "expression": "equipment_code != null",
                    "severity": "blocking",
                    "natural_language": "实验设备必须具备资产编号。",
                }
            ],
            "tableKeywords": ["equipment"],
            "capabilityTags": ["asset-governance"],
        },
    )
    assert blueprint["source"] == "custom"
    assert any(item["id"] == "lab-asset" for item in list_industry_blueprints(platform_db))
    assert infer_industry_blueprint(["equipment"], platform_db=platform_db).id == "lab-asset"

    source = register_data_source(platform_db, "实验室资产系统", "sqlite", str(legacy_db), domain="实验室资产")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="lab-asset")

    assert ontology["blueprint"]["id"] == "lab-asset"
    assert any(item["code"] == "equipment" and item["name"] == "实验设备" for item in ontology["objects"])
    assert any(item["code"] == "equipment_code" and item["name"] == "资产编号" for item in ontology["attributes"])
    assert any(rule["code"] == "lab_asset_code_required" for rule in ontology["rules"])


def test_mapping_review_and_publish_governance_flow(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id)

    mappings = list_semantic_mappings(platform_db, ontology["id"])
    assert mappings["items"]
    assert {item["status"] for item in mappings["items"]} == {"pending"}

    initial_release = assess_ontology_release_readiness(platform_db, ontology["id"])
    assert initial_release["status"] == "blocked"
    assert any(item["code"] == "pending_mappings" and not item["passed"] for item in initial_release["gates"])

    first_mapping_id = mappings["items"][0]["id"]
    reviewed = review_semantic_mapping(platform_db, first_mapping_id, "confirmed", "业务专家", "表对象映射正确")
    assert reviewed["status"] == "confirmed"
    assert reviewed["reviewer"] == "业务专家"

    try:
        publish_ontology(platform_db, ontology["id"], "架构师")
    except ValueError as error:
        assert "待审核" in str(error)
    else:
        raise AssertionError("存在待审核映射时不应发布本体")

    bulk = bulk_review_semantic_mappings(platform_db, ontology["id"], "confirmed", "业务专家", "MVP 批量确认")
    assert bulk["reviewedCount"] == len(mappings["items"]) - 1

    reviewed_release = assess_ontology_release_readiness(platform_db, ontology["id"])
    assert reviewed_release["status"] in {"ready", "review"}
    assert reviewed_release["summary"]["confirmedMappings"] == len(mappings["items"])
    assert reviewed_release["summary"]["blockers"] == 0
    assert any(item["code"] == "data_source_1_schema_drift" and item["passed"] for item in reviewed_release["gates"])

    custom_rule = upsert_business_rule(
        platform_db,
        ontology["id"],
        "contract_title_required",
        "合同标题不能为空",
        "validation",
        "contract",
        "title != null",
        "blocking",
        "合同标题不能为空。",
        "规则管理员",
    )
    assert custom_rule["code"] == "contract_title_required"
    assert any(
        rule["code"] == "contract_title_required" for rule in list_business_rules(platform_db, ontology["id"])["items"]
    )

    published = publish_ontology(platform_db, ontology["id"], "架构师")
    assert published["status"] == "published"
    assert published["mappingCounts"]["confirmed"] == len(mappings["items"])

    try:
        generate_ontology_draft(platform_db, source.id)
    except ValueError as error:
        assert "本体版本已发布" in str(error)
    else:
        raise AssertionError("已发布本体版本不应被重新生成草案覆盖")

    try:
        review_semantic_mapping(platform_db, first_mapping_id, "rejected", "业务专家", "发布后尝试修改")
    except ValueError as error:
        assert "已发布本体" in str(error)
    else:
        raise AssertionError("发布后的语义映射不应允许直接修改")

    try:
        upsert_business_rule(
            platform_db,
            ontology["id"],
            "contract_title_required",
            "合同标题不能为空",
            "validation",
            "contract",
            "title != null",
            "blocking",
            "合同标题不能为空。",
            "规则管理员",
        )
    except ValueError as error:
        assert "已发布本体" in str(error)
    else:
        raise AssertionError("发布后的规则不应允许直接修改")

    derived = derive_ontology_version(platform_db, ontology["id"], "0.2.0", "架构师")
    assert derived["status"] == "draft"
    assert derived["version"] == "0.2.0"
    assert derived["sourceOntologyId"] == ontology["id"]
    assert derived["mappingCount"] == len(mappings["items"])

    derived_mappings = list_semantic_mappings(platform_db, derived["id"])
    assert {item["status"] for item in derived_mappings["items"]} == {"pending"}
    assert all("需重新审核" in item["evidence"] for item in derived_mappings["items"])

    derived_rule = upsert_business_rule(
        platform_db,
        derived["id"],
        "contract_title_required",
        "合同标题不能为空",
        "validation",
        "contract",
        "title != null",
        "blocking",
        "合同标题不能为空。",
        "规则管理员",
    )
    assert derived_rule["code"] == "contract_title_required"

    try:
        derive_ontology_version(platform_db, derived["id"], "0.3.0", "架构师")
    except ValueError as error:
        assert "已发布本体" in str(error)
    else:
        raise AssertionError("不应从未发布草案派生新版本")


def test_scan_records_column_profile(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)

    with connect(platform_db) as conn:
        status = conn.execute(
            """
            select sc.*
            from source_column sc
            join source_table st on st.id = sc.source_table_id
            where st.table_name = 'contract' and sc.column_name = 'status'
            """
        ).fetchone()

    assert status is not None
    assert status["enum_candidate"] == 1
    assert status["distinct_count"] == 3


def test_schema_drift_analysis_compares_live_database_with_scanned_baseline(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")

    not_scanned = analyze_schema_drift(platform_db, source.id)
    assert not_scanned["status"] == "not_scanned"
    assert not_scanned["nextActions"]

    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id, blueprint_id="contract-management")

    stable = analyze_schema_drift(platform_db, source.id)
    assert stable["status"] == "no_drift"
    assert stable["summary"]["changedTables"] == 0

    with sqlite3.connect(legacy_db) as conn:
        conn.execute("alter table contract add column risk_level text")
        conn.execute(
            """
            create table contract_attachment (
                id integer primary key,
                contract_id integer not null,
                file_name text not null,
                foreign key(contract_id) references contract(id)
            )
            """
        )

    drift = analyze_schema_drift(platform_db, source.id)

    assert drift["status"] == "drift_detected"
    assert drift["summary"]["addedTables"] == 1
    assert drift["summary"]["changedTables"] == 1
    assert drift["summary"]["addedColumns"] == 1
    assert drift["summary"]["impactedObjects"] >= 1
    assert {item["tableName"] for item in drift["addedTables"]} == {"contract_attachment"}
    contract_change = next(item for item in drift["changedTables"] if item["tableName"] == "contract")
    assert [item["columnName"] for item in contract_change["addedColumns"]] == ["risk_level"]
    assert any(
        item["ontologyId"] == ontology["id"] and item["code"] == "contract" for item in drift["impacts"]["objects"]
    )
    assert any("重新扫描" in action or "本体" in action for action in drift["nextActions"])


def test_data_source_connection_checks_sqlite_file(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")

    registered = check_data_source_connection(platform_db, data_source_id=source.id)
    missing = check_data_source_connection(
        platform_db, source_type="sqlite", connection_uri=str(tmp_path / "missing.sqlite3")
    )

    assert registered["reachable"] is True
    assert registered["status"] == "ok"
    assert missing["reachable"] is False
    assert missing["status"] == "not_found"


def test_business_api_gateway_check_reports_configuration_and_reachability(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    class GatewayHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            self.send_response(204)
            self.end_headers()

        def log_message(self, format: str, *args: object) -> None:
            return None

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(
        platform_db, "合同管理 API 系统", "sqlite", str(legacy_db), domain="合同管理", system_category="database+api"
    )

    missing = check_business_api_gateway(platform_db, source.id)
    assert missing["configured"] is False
    assert missing["reachable"] is False
    assert missing["status"] == "not_configured"

    invalid = register_data_source(
        platform_db,
        "合同管理 API 系统",
        "sqlite",
        str(legacy_db),
        domain="合同管理",
        system_category="database+api",
        api_base_url="ftp://legacy.example/api",
    )
    invalid_result = check_business_api_gateway(platform_db, invalid.id)
    assert invalid_result["configured"] is True
    assert invalid_result["reachable"] is False
    assert invalid_result["status"] == "invalid_url"

    server = HTTPServer(("127.0.0.1", 0), GatewayHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        api_base_url = f"http://127.0.0.1:{server.server_port}/"
        reachable = register_data_source(
            platform_db,
            "合同管理 API 系统",
            "sqlite",
            str(legacy_db),
            domain="合同管理",
            system_category="database+api",
            api_base_url=api_base_url,
        )
        result = check_business_api_gateway(platform_db, reachable.id)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)

    assert result["configured"] is True
    assert result["reachable"] is True
    assert result["status"] == "ok"
    assert result["statusCode"] == 204
    assert result["apiBaseUrl"] == api_base_url


def test_platform_lists_data_sources_and_ontologies_for_control_plane(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id)

    data_sources = list_data_sources(platform_db)
    ontologies = list_ontologies(platform_db)

    assert len(data_sources) == 1
    assert data_sources[0]["id"] == source.id
    assert data_sources[0]["name"] == "合同管理样例系统"
    assert data_sources[0]["domain"] == "合同管理"
    assert data_sources[0]["systemCategory"] == "database"
    assert data_sources[0]["system_category"] == "database"
    assert data_sources[0]["sourceType"] == "sqlite"
    assert data_sources[0]["source_type"] == "sqlite"
    assert data_sources[0]["connectionUri"] == str(legacy_db)
    assert data_sources[0]["connection_uri"] == str(legacy_db)
    assert data_sources[0]["apiBaseUrl"] == ""
    assert data_sources[0]["api_base_url"] == ""
    assert data_sources[0]["apiHeadersConfigured"] is False
    assert data_sources[0]["apiHeaderNames"] == []
    assert data_sources[0]["capabilities"] == ["metadata_scan", "semantic_mapping"]
    assert data_sources[0]["createdAt"]
    assert ontologies[0]["id"] == ontology["id"]
    assert ontologies[0]["name"] == ontology["name"]
    assert ontologies[0]["domain"] == "合同管理"
    assert ontologies[0]["version"] == "0.1.0"
    assert ontologies[0]["status"] == "draft"


def test_ontology_exports_jsonld_and_turtle_assets(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id)

    jsonld = export_ontology_asset(platform_db, ontology["id"], "jsonld")
    document = json.loads(jsonld["content"])
    graph_types = {node["@type"] for node in document["@graph"]}

    assert jsonld["mediaType"] == "application/ld+json"
    assert jsonld["filename"].endswith(".jsonld")
    assert "ont:Ontology" in graph_types
    assert "ont:BusinessObject" in graph_types
    assert "ont:BusinessAttribute" in graph_types
    assert "ont:BusinessRule" in graph_types
    assert any(node.get("code") == "contract" for node in document["@graph"])

    turtle = export_ontology_asset(platform_db, ontology["id"], "turtle")
    assert turtle["mediaType"] == "text/turtle"
    assert turtle["filename"].endswith(".ttl")
    assert "@prefix ont:" in turtle["content"]
    assert "ont:BusinessObject" in turtle["content"]

    # Parsed, not pattern-matched. The previous version of this test asserted the output
    # contained `bp:object/contract` -- which is *itself* invalid Turtle, because a
    # prefixed name's local part cannot contain `/`. So the assertion pinned the bug in
    # place: the export shipped unparseable by any RDF tool, and this test stayed green,
    # because a syntactically invalid document contains a substring just as well as a
    # valid one.
    _assert_parses_as_turtle(turtle["content"])

    try:
        export_ontology_asset(platform_db, ontology["id"], "xml")
    except ValueError as error:
        assert "jsonld" in str(error)
    else:
        raise AssertionError("不支持的本体导出格式应被拒绝")


def _assert_parses_as_turtle(content: str) -> None:
    """Parse with a real RDF library, or state plainly that it could not be checked.

    `rdflib` is not a dependency of this project -- the kernel is dependency-free by
    design -- so this skips rather than fails when it is absent. Skipping is honest;
    silently passing would restore exactly the false confidence this test exists to
    remove. CI installs it, so the check does run where it gates merges.
    """
    rdflib = pytest.importorskip("rdflib", reason="需要 rdflib 才能验证 Turtle 语法")
    graph = rdflib.Graph()
    graph.parse(data=content, format="turtle")
    assert len(graph) > 0, "Turtle 解析成功但没有三元组"


def test_model_config_persists_masks_and_resets(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    initial = get_model_config(platform_db)
    assert initial["provider"] == "openrouter"
    assert initial["hasApiKey"] is False

    saved = update_model_config(
        platform_db,
        {
            "apiKey": "or-test-secret-key",
            "model": "openai/test-model",
            "baseUrl": "https://openrouter.ai/api/v1/",
            "httpReferer": "https://example.com/app",
            "appTitle": "Ontology Platform Test",
            "serviceTier": "auto",
            "timeoutSeconds": 12,
        },
    )
    assert saved["success"] is True

    persisted = get_model_config(platform_db)
    assert persisted["configured"] is True
    assert persisted["hasApiKey"] is True
    assert persisted["apiKey"] == "or-tes...-key"
    assert persisted["model"] == "openai/test-model"
    assert persisted["baseUrl"] == "https://openrouter.ai/api/v1"
    assert persisted["httpReferer"] == "https://example.com/app"
    assert persisted["appTitle"] == "Ontology Platform Test"
    assert persisted["timeoutSeconds"] == 12.0

    update_model_config(platform_db, {"apiKey": "", "model": "openai/another-model"})
    retained_key = get_model_config(platform_db)
    assert retained_key["apiKey"] == "or-tes...-key"
    assert retained_key["model"] == "openai/another-model"

    reset = reset_model_config(platform_db)
    assert reset["success"] is True
    assert get_model_config(platform_db)["hasApiKey"] is False


def test_legacy_model_config_table_migrates_to_single_row_schema(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    with connect(platform_db) as conn:
        conn.execute("drop table model_config")
        conn.execute(
            """
            create table model_config (
                id integer primary key autoincrement,
                config_key text not null unique,
                config_value text not null,
                description text not null default '',
                updated_at text not null default current_timestamp
            )
            """
        )
        conn.executemany(
            "insert into model_config (config_key, config_value) values (?, ?)",
            [
                ("api_key", "or-legacy-secret-key"),
                ("model", "deepseek/deepseek-v4-flash"),
                ("base_url", "https://openrouter.ai/api/v1"),
                ("http_referer", "https://legacy.example"),
                ("app_title", "Legacy Ontology Platform"),
                ("service_tier", "auto"),
                ("timeout_seconds", "15"),
            ],
        )

    initialize_platform_db(platform_db)
    migrated = get_model_config(platform_db)

    assert migrated["hasApiKey"] is True
    assert migrated["apiKey"] == "or-leg...-key"
    assert migrated["model"] == "deepseek/deepseek-v4-flash"
    assert migrated["httpReferer"] == "https://legacy.example"
    assert migrated["appTitle"] == "Legacy Ontology Platform"
    assert migrated["timeoutSeconds"] == 15.0


def test_database_adapter_registry_supports_common_enterprise_databases(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    sqlite_adapter = get_adapter("sqlite")
    postgres_adapter = get_adapter("postgresql")
    mysql_adapter = get_adapter("mysql")

    assert sqlite_adapter.source_type == "sqlite"
    assert postgres_adapter.source_type == "postgresql"
    assert mysql_adapter.source_type == "mysql"

    postgres = register_data_source(
        platform_db,
        "PostgreSQL 生产系统占位",
        "postgresql",
        "postgresql://user:pass@127.0.0.1:5432/app",
        domain="通用业务",
    )
    mysql = register_data_source(
        platform_db,
        "MySQL 生产系统占位",
        "mysql",
        "mysql://user:pass@127.0.0.1:3306/app",
        domain="通用业务",
    )

    assert postgres.source_type == "postgresql"
    assert mysql.source_type == "mysql"


def test_postgresql_scan_reports_missing_driver_as_actionable_error(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    source = register_data_source(
        platform_db,
        "PostgreSQL 生产系统占位",
        "postgresql",
        "postgresql://user:pass@127.0.0.1:5432/app",
        domain="通用业务",
    )

    try:
        scan_data_source(platform_db, source.id)
    except ValueError as error:
        message = str(error)
        # Either the driver is missing and the message names it, or the server answered and
        # the message names what to check. The `or "connection"` this replaces matched the
        # driver's own English wording -- which is exactly what the message layer stopped
        # showing, so it failed on the one machine where psycopg *is* installed (CI) while
        # passing everywhere it is not.
        assert "psycopg" in message or any(
            hint in message for hint in ("认证失败", "无法连接", "数据库不存在", "没有读取权限")
        ), message
        # And it is a message, not a driver exception that leaked through.
        assert "Traceback" not in message
        assert "psycopg." not in message and "OperationalError" not in message
    else:
        raise AssertionError("测试环境没有真实 PostgreSQL 服务时不应扫描成功")


def test_postgresql_connection_test_reports_actionable_status(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"

    initialize_platform_db(platform_db)
    result = check_data_source_connection(
        platform_db,
        source_type="postgresql",
        connection_uri="postgresql://user:pass@127.0.0.1:1/app",
    )

    assert result["reachable"] is False
    assert result["status"] in {"driver_missing", "connection_error"}
    assert result["message"]


def test_equipment_domain_generates_industry_rules_and_assessment(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_equipment.sqlite3"

    initialize_platform_db(platform_db)
    create_equipment_sample_db(legacy_db)
    source = register_data_source(
        platform_db, "设备运维样例系统", "sqlite", str(legacy_db), domain="设备运维", system_category="database+api"
    )
    api = register_source_api(
        platform_db, source.id, "close_work_order", "关闭工单", "POST", "/work-orders/{id}/close", "work_order.close"
    )
    scan = scan_data_source(platform_db, source.id)

    assert len(scan["tables"]) == 4
    assert api.operation_code == "close_work_order"
    assert list_source_apis(platform_db, source.id)[0]["semantic_action"] == "work_order.close"

    ontology = generate_ontology_draft(platform_db, source.id)
    object_codes = {item["code"] for item in ontology["objects"]}
    assert {"equipment", "work_order", "inspection_record", "spare_part"}.issubset(object_codes)
    assert any(rule["code"] == "critical_equipment_open_fault" for rule in ontology["rules"])

    assessment = assess_instance(platform_db, ontology["id"], "equipment", "1")
    assert assessment["decision"]["status"] == "review"
    failed_rules = {item["ruleCode"] for item in assessment["ruleResults"] if not item["passed"]}
    assert "critical_equipment_open_fault" in failed_rules


def test_openrouter_client_uses_subscription_key_shape() -> None:
    calls: list[dict[str, object]] = []

    def fake_transport(
        url: str, headers: dict[str, str], payload: dict[str, object], timeout_seconds: float
    ) -> dict[str, object]:
        calls.append({"url": url, "headers": headers, "payload": payload, "timeout": timeout_seconds})
        return {
            "model": "openai/test-model",
            "choices": [{"message": {"content": "ok"}}],
            "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
        }

    client = OpenRouterClient(
        OpenRouterConfig(
            api_key="or-test-key",
            model="openai/test-model",
            base_url="https://openrouter.ai/api/v1",
            http_referer="https://example.com",
            app_title="Ontology Platform",
            service_tier="auto",
            timeout_seconds=10,
        ),
        transport=fake_transport,
    )

    result = client.chat([{"role": "user", "content": "ping"}], session_id="session-1")

    assert result.content == "ok"
    assert calls[0]["url"] == "https://openrouter.ai/api/v1/chat/completions"
    headers = calls[0]["headers"]
    payload = calls[0]["payload"]
    assert headers["Authorization"] == "Bearer or-test-key"
    assert headers["HTTP-Referer"] == "https://example.com"
    # OpenRouter reads X-Title. The previous X-OpenRouter-Title was not a header
    # any provider recognises, so attribution was silently dropped.
    assert headers["X-Title"] == "Ontology Platform"
    assert payload["model"] == "openai/test-model"
    assert payload["service_tier"] == "auto"
    assert payload["session_id"] == "session-1"


def test_blueprint_draft_uses_local_fallback_and_openrouter_shape(tmp_path: Path) -> None:
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"

    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)

    local = generate_blueprint_draft(
        platform_db,
        source.id,
        OpenRouterClient(
            OpenRouterConfig("", "test", "https://openrouter.ai/api/v1", "", "Ontology Platform", "auto", 10)
        ),
    )
    assert local["usedRemoteModel"] is False
    assert local["blueprint"]["domain"] == "合同管理"
    assert "contract" in local["blueprint"]["objectHints"]

    calls: list[dict[str, object]] = []

    def fake_transport(
        url: str, headers: dict[str, str], payload: dict[str, object], timeout_seconds: float
    ) -> dict[str, object]:
        calls.append({"url": url, "headers": headers, "payload": payload, "timeout": timeout_seconds})
        return {
            "model": "openai/test-model",
            "choices": [
                {
                    "message": {
                        "content": json.dumps(
                            {
                                "id": 123,
                                "name": "AI 合同蓝图",
                                "domain": "合同管理",
                                "description": "AI 生成",
                                "objectHints": {"contract": "合同"},
                                "attributeHints": {"contract.contract_no": "合同编号"},
                                "rules": [],
                                "tableKeywords": ["contract"],
                                "capabilityTags": ["ai-generated"],
                            },
                            ensure_ascii=False,
                        )
                    }
                }
            ],
            "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
        }

    remote = generate_blueprint_draft(
        platform_db,
        source.id,
        OpenRouterClient(
            OpenRouterConfig(
                "or-test-key", "openai/test-model", "https://openrouter.ai/api/v1", "", "Ontology Platform", "auto", 10
            ),
            transport=fake_transport,
        ),
    )
    assert remote["usedRemoteModel"] is True
    assert remote["blueprint"]["id"] == "123"
    assert remote["blueprint"]["attributeHints"]["contract_no"] == "合同编号"
    assert calls[0]["payload"]["service_tier"] == "auto"


def test_api_module_imports_on_supported_python() -> None:
    api = importlib.import_module("ontology_platform.api")
    assert api.app.title == "本体改造研发平台"


def _blueprint_client(blueprint_payload: dict, calls: list, model: str = "openai/test-model") -> OpenRouterClient:
    """A client whose transport answers with `blueprint_payload`, recording the request."""

    def fake_transport(
        url: str, headers: dict[str, str], payload: dict[str, object], timeout_seconds: float
    ) -> dict[str, object]:
        calls.append({"url": url, "headers": headers, "payload": payload})
        return {
            "model": model,
            "choices": [{"message": {"content": json.dumps(blueprint_payload, ensure_ascii=False)}}],
            "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
        }

    return OpenRouterClient(
        OpenRouterConfig("or-test-key", model, "https://openrouter.ai/api/v1", "", "Ontology Platform", "auto", 10),
        transport=fake_transport,
    )


def _scanned_source(tmp_path: Path, name: str = "platform.sqlite3"):
    platform_db = tmp_path / name
    legacy_db = tmp_path / f"{name}.legacy.sqlite3"
    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    return platform_db, source


def test_blueprint_hints_written_as_a_list_of_objects_are_read(tmp_path: Path) -> None:
    """The shape a live model actually returned, which used to discard the whole response.

    The prompt named `objectHints` without stating its shape, and a real model answered
    with a list of object descriptors -- a reasonable reading, and the richer one, since it
    carries the source table and key that a `{code: label}` mapping cannot. The normalizer
    ran `dict()` over that list, raised, and the caller's broad `except` replaced sixteen
    kilobytes of usable modelling with the local heuristic, reporting only
    `usedRemoteModel: False` -- which reads as "no model was configured".
    """
    platform_db, source = _scanned_source(tmp_path)
    calls: list[dict] = []
    draft = generate_blueprint_draft(
        platform_db,
        source.id,
        _blueprint_client(
            {
                "id": "contract-management-blueprint",
                "name": "合同管理行业蓝图",
                "domain": "合同管理",
                "description": "模型生成",
                "objectHints": [
                    {
                        "code": "Contract",
                        "name": "合同",
                        "sourceTable": "contract",
                        "primaryKey": "id",
                        "businessKey": ["contract_no"],
                        "relationships": [{"type": "many-to-one", "targetObjectCode": "Customer"}],
                    },
                    {"code": "Customer", "name": "客户", "sourceTable": "customer"},
                ],
                "attributeHints": [
                    {"code": "contract.amount", "name": "合同金额"},
                    {"code": "contract.contract_no", "name": "合同编号"},
                ],
                "rules": [
                    {
                        "code": "amount_positive",
                        "name": "金额为正",
                        "rule_type": "validation",
                        "scope_object_code": "Contract",
                        "expression": "amount > 0",
                        "severity": "blocking",
                        "natural_language": "合同金额必须大于零。",
                    }
                ],
                "tableKeywords": ["contract", "customer"],
                "capabilityTags": ["contract_lifecycle"],
            },
            calls,
        ),
    )

    assert draft["usedRemoteModel"] is True, draft.get("remoteError")
    assert "unreadableModelFields" not in draft
    blueprint = draft["blueprint"]
    # The model's business names survive, rather than the snake_case-derived heuristic ones.
    assert blueprint["objectHints"] == {"Contract": "合同", "Customer": "客户"}
    assert blueprint["attributeHints"]["contract.amount"] == "合同金额"
    # Attribute hints keep the dotted spelling and gain the bare column, as before.
    assert blueprint["attributeHints"]["amount"] == "合同金额"
    assert [rule["code"] for rule in blueprint["rules"]] == ["amount_positive"]
    assert calls, "模型未被调用"


def test_a_bare_string_keyword_list_is_not_exploded_into_characters(tmp_path: Path) -> None:
    """Iterating a string yields its characters: "contract" became six keywords."""
    platform_db, source = _scanned_source(tmp_path)
    draft = generate_blueprint_draft(
        platform_db,
        source.id,
        _blueprint_client(
            {
                "id": "x",
                "name": "n",
                "domain": "合同管理",
                "objectHints": {"contract": "合同"},
                "attributeHints": {},
                "rules": [],
                "tableKeywords": "contract",
                "capabilityTags": "contract_lifecycle",
            },
            [],
        ),
    )

    assert draft["blueprint"]["tableKeywords"] == ["contract"]
    assert draft["blueprint"]["capabilityTags"] == ["contract_lifecycle"]


def test_one_unreadable_field_costs_one_field(tmp_path: Path) -> None:
    """The model's other work must survive a single shape it got wrong.

    Before this, any shape surprise raised out of the normalizer and the whole response was
    discarded. The failure has to be proportionate to the mistake, and it has to be visible:
    the caller asked for a model-assisted draft and is entitled to know which parts of it
    fell back.
    """
    platform_db, source = _scanned_source(tmp_path)
    draft = generate_blueprint_draft(
        platform_db,
        source.id,
        _blueprint_client(
            {
                "id": "partial",
                "name": "部分可用",
                "domain": "合同管理",
                "description": "模型生成",
                "objectHints": "请把合同和客户建模为业务对象",
                "attributeHints": {},
                "rules": [],
                "tableKeywords": ["contract"],
                "capabilityTags": [],
            },
            [],
        ),
    )

    assert draft["usedRemoteModel"] is True
    assert draft["unreadableModelFields"] == ["objectHints"]
    # The fields the model did supply are used...
    assert draft["blueprint"]["name"] == "部分可用"
    assert draft["blueprint"]["tableKeywords"] == ["contract"]
    # ...and the unreadable one falls back to the draft's own hints rather than to empty.
    assert "contract" in draft["blueprint"]["objectHints"]


class TestAModelAnsweringWithSentinels:
    """The understanding prompt asks the model to write `null` for anything it cannot
    determine, and models write that as the string "null" at least as often as JSON null.

    These drive the real `query_natural_language` with a fake transport, because the defect
    they pin was only reachable through the model's answer: the values went straight from
    parsed JSON into `int()`, and the resulting `ValueError` travelled out to the asker as
    HTTP 400 reading `invalid literal for int() with base 10: 'null'`.
    """

    @staticmethod
    def _model_returning(payload: dict):
        def fake_transport(
            url: str, headers: dict[str, str], request_payload: dict[str, object], timeout_seconds: float
        ) -> dict[str, object]:
            return {
                "model": "openai/test-model",
                "choices": [{"message": {"content": json.dumps(payload, ensure_ascii=False)}}],
                "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
            }

        return fake_transport

    def _scenario(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, model_payload: dict):
        import ontology_platform.model_client as model_client_module

        platform_db = tmp_path / "platform.sqlite3"
        legacy_db = tmp_path / "legacy.sqlite3"
        initialize_platform_db(platform_db)
        create_contract_sample_db(legacy_db)
        source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
        scan_data_source(platform_db, source.id)
        ontology = generate_ontology_draft(platform_db, source.id)

        monkeypatch.setattr(model_client_module, "_post_json", self._model_returning(model_payload))
        update_model_config(platform_db, {"apiKey": "or-test-key", "model": "openai/test-model"})
        return platform_db, source, ontology

    @pytest.mark.parametrize("sentinel", ["null", "unknown", "N/A"])
    def test_string_sentinels_in_identifier_fields_are_treated_as_absent(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, sentinel: str
    ) -> None:
        platform_db, source, ontology = self._scenario(
            tmp_path,
            monkeypatch,
            {
                "intent": "compliance_assessment",
                "ontologyId": sentinel,
                "dataSourceId": sentinel,
                "objectCode": sentinel,
                "instanceId": sentinel,
                "operationCode": None,
            },
        )

        # The question names nothing, so the platform resolves the target itself. What
        # matters is that it answers at all rather than surfacing a parse error.
        answer = query_natural_language(
            platform_db,
            "合同是否合规？",
            ontology["ontology"]["id"],
            source.id,
            use_model=True,
        )

        assert answer["answer"], answer
        assert answer["resolved"]["ontologyId"] == ontology["ontology"]["id"]

    def test_a_structured_identifier_does_not_become_an_object_code(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """`str({"code": "x"})` is truthy, so this used to be searched for literally."""
        platform_db, source, ontology = self._scenario(
            tmp_path,
            monkeypatch,
            {
                "intent": "compliance_assessment",
                "ontologyId": {"value": 1},
                "dataSourceId": None,
                "objectCode": {"code": "Contract"},
                "instanceId": ["1"],
                "operationCode": None,
            },
        )

        answer = query_natural_language(
            platform_db,
            "合同是否合规？",
            ontology["ontology"]["id"],
            source.id,
            use_model=True,
        )

        assert answer["answer"], answer
        # The structured values are dropped, so resolution falls back rather than trying to
        # find an object literally called "{'code': 'Contract'}".
        assert "{'code'" not in json.dumps(answer, ensure_ascii=False)

    def test_an_intent_outside_the_enumerated_set_does_not_suppress_routing(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A near miss like "compliance" is not an interpretation of the question.

        Letting it through routed the question to no branch, so the asker got the
        unknown-intent answer for a question the deterministic detector routes fine.
        """
        platform_db, source, ontology = self._scenario(
            tmp_path,
            monkeypatch,
            {
                "intent": "compliance",
                "ontologyId": None,
                "dataSourceId": None,
                "objectCode": None,
                "instanceId": None,
                "operationCode": None,
            },
        )

        answer = query_natural_language(
            platform_db,
            "合同是否合规？",
            ontology["ontology"]["id"],
            source.id,
            use_model=True,
        )

        assert answer["intent"] == "compliance_assessment", answer


class TestBlueprintImportShapes:
    """Importing a blueprint, which is where a generated draft has to come back in.

    The platform writes a blueprint to the database through this function, so whatever
    `generate_blueprint_draft` can produce, this has to be able to accept -- otherwise the
    platform generates a draft it then refuses to import. It could not accept the shape a
    live model produced, and it accepted a neighbouring one *silently wrong*.
    """

    def test_a_descriptor_list_is_read_faithfully(self, tmp_path: Path) -> None:
        """The live-model shape: a list of objects rather than a mapping."""
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)

        stored = upsert_industry_blueprint(
            platform_db,
            {
                "id": "ai-contract",
                "name": "AI 合同蓝图",
                "domain": "合同管理",
                "objectHints": [
                    {"code": "contract", "name": "合同", "sourceTable": "contract", "primaryKey": "id"},
                    {"code": "customer", "name": "客户"},
                ],
                "attributeHints": [{"code": "contract.amount", "name": "合同金额"}],
            },
        )

        assert stored["objectHints"] == {"contract": "合同", "customer": "客户"}
        assert stored["attributeHints"]["contract.amount"] == "合同金额"
        # Attribute hints gain the bare column alongside the dotted key, as before.
        assert stored["attributeHints"]["amount"] == "合同金额"

    def test_two_key_descriptors_are_not_read_as_a_key_value_pair(self, tmp_path: Path) -> None:
        """The silent case, which is worse than the crash it sat next to.

        `dict([{"column": "amount", "label": "合同金额"}])` does not raise: a two-element
        dict is itself a valid key/value pair, so the *keys* became the entry. The platform
        stored an object named "column" labelled "label", reported success, and nothing
        anywhere recorded that the model's answer had been misread.
        """
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)

        stored = upsert_industry_blueprint(
            platform_db,
            {
                "id": "ai",
                "name": "n",
                "domain": "d",
                "objectHints": [{"column": "amount", "label": "合同金额"}, {"column": "status", "label": "状态"}],
            },
        )

        assert stored["objectHints"] == {"amount": "合同金额", "status": "状态"}
        assert "column" not in stored["objectHints"]

    def test_the_documented_mapping_form_is_unchanged(self, tmp_path: Path) -> None:
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)
        stored = upsert_industry_blueprint(
            platform_db,
            {"id": "ai", "name": "n", "domain": "d", "objectHints": {"contract": "合同"}},
        )
        assert stored["objectHints"] == {"contract": "合同"}

    def test_a_list_of_candidate_labels_is_not_stored_as_a_repr(self, tmp_path: Path) -> None:
        """`str(["合同", "销售合同"])` is truthy, so it used to be accepted as the label.

        Two candidate names cannot be resolved without inventing a rule, so the code is
        used as the label instead -- a name the model did write, rather than a Python repr
        of names it offered.
        """
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)
        stored = upsert_industry_blueprint(
            platform_db,
            {"id": "ai", "name": "n", "domain": "d", "objectHints": {"contract": ["合同", "销售合同"]}},
        )
        assert stored["objectHints"] == {"contract": "contract"}

    def test_a_generated_draft_imports_back(self, tmp_path: Path) -> None:
        """The closed loop, which is the property that was broken.

        Nothing linked the drafting path to the import path, so each was tested against its
        own idea of the shape and the pair did not fit.
        """
        platform_db, source = _scanned_source(tmp_path)
        draft = generate_blueprint_draft(
            platform_db,
            source.id,
            _blueprint_client(
                {
                    "id": "loop",
                    "name": "回路蓝图",
                    "domain": "合同管理",
                    "description": "模型生成",
                    "objectHints": [{"code": "contract", "name": "合同", "sourceTable": "contract"}],
                    "attributeHints": [{"code": "contract.amount", "name": "合同金额"}],
                    "rules": [],
                    "tableKeywords": ["contract"],
                    "capabilityTags": ["modelled"],
                },
                [],
            ),
        )
        assert draft["usedRemoteModel"] is True

        stored = upsert_industry_blueprint(platform_db, dict(draft["blueprint"]))

        assert stored["objectHints"] == {"contract": "合同"}
        assert stored["attributeHints"]["amount"] == "合同金额"


class TestRuleElementsAreReadableOrRefused:
    """A blueprint's rules, which the import path reads by subscripting.

    `item["code"]` raised `KeyError` for an element missing a key, and the route converts
    only `ValueError` into a response -- so a user pasting a model-generated draft got a 500
    with no indication of which rule was at fault. The quieter half was worse: a missing
    optional field became the *literal string* `"None"` via `str(item.get(...))`, and a rule
    whose scope is `"None"` matches no object, so `ontology._create_blueprint_rules` dropped
    it without a word. The blueprint then declares more rules than it creates.
    """

    @pytest.mark.parametrize(
        "rules, expected",
        [
            ([{"code": "r1", "name": "n", "rule_type": "validation"}], "expression"),
            ([{"code": "r2", "name": "n", "expression": "a > 0", "severity": "blocking"}], "scope_object_code"),
            (["合同金额必须大于 0"], "不是对象"),
            ({"r1": {"name": "n"}}, "不是对象"),
            ([None], "不是对象"),
        ],
    )
    def test_an_unreadable_rule_is_refused_with_a_reason(self, tmp_path: Path, rules: object, expected: str) -> None:
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)

        with pytest.raises(ValueError) as refused:
            upsert_industry_blueprint(platform_db, {"id": "x", "name": "n", "domain": "d", "rules": rules})

        assert expected in str(refused.value), refused.value

    def test_a_missing_optional_field_is_empty_not_the_string_none(self, tmp_path: Path) -> None:
        platform_db = tmp_path / "platform.sqlite3"
        initialize_platform_db(platform_db)

        stored = upsert_industry_blueprint(
            platform_db,
            {
                "id": "x",
                "name": "n",
                "domain": "d",
                "rules": [
                    {
                        "code": "r3",
                        "name": "n",
                        "scope_object_code": "contract",
                        "expression": "a > 0",
                        "severity": "blocking",
                    }
                ],
            },
        )

        rule = stored["rules"][0]
        assert rule["rule_type"] == ""
        assert rule["natural_language"] == ""
        assert "None" not in json.dumps(rule)


class TestTheReasoningChainNormalizer:
    """The sibling of the blueprint normalizer, which had no test at all.

    The blueprint's prompt and normalizer were hardened while this one kept validating only
    `isinstance(list)` -- so a model answering per its own reading ("给出两种路径" -> one
    object per path) produced a chain the UI could not draw: React rejects an object child
    and the error boundary replaces the page, with the endpoint having returned 200 and
    recorded success.
    """

    @staticmethod
    def _local() -> dict:
        return {
            "summary": "本地摘要",
            "reasoningSteps": ["本地步骤"],
            "proposedObjects": [{"objectCode": "t", "objectName": "T"}],
            "proposedRelations": [],
            "proposedRules": [],
            "buildPlan": ["本地路径"],
            "questionsForUser": ["本地问题"],
        }

    def _normalize(self, parsed: dict) -> tuple[dict, list[str]]:
        from ontology_platform.model_client import _normalize_reasoning_chain

        return _normalize_reasoning_chain(parsed, self._local())

    def test_objects_where_text_is_expected_keep_the_models_words(self) -> None:
        """The shape the prompt's own wording invites."""
        chain, unreadable = self._normalize(
            {
                "summary": "模型摘要",
                "reasoningSteps": [{"step": 1, "detail": "读取表与字段"}],
                "buildPlan": [{"path": "自动构建", "steps": ["生成草案"]}],
            }
        )

        assert chain["reasoningSteps"] == ["读取表与字段"]
        assert chain["buildPlan"] == ["自动构建"]
        assert unreadable == []

    def test_strings_where_records_are_expected_fall_back_and_say_so(self) -> None:
        """Otherwise every row renders "-" and the candidate list looks empty."""
        chain, unreadable = self._normalize({"proposedObjects": ["contract", "customer"]})

        assert unreadable == ["proposedObjects"]
        assert chain["proposedObjects"] == self._local()["proposedObjects"]

    def test_a_summary_written_as_lines_is_joined(self) -> None:
        chain, unreadable = self._normalize({"summary": ["第一点", "第二点"]})

        assert chain["summary"] == "第一点\n第二点"
        assert unreadable == []

    def test_every_element_is_a_string_or_a_record_as_the_frontend_expects(self) -> None:
        chain, _ = self._normalize(
            {
                "reasoningSteps": [{"detail": "一"}, "二", None, {"unknown": "shape"}],
                "proposedObjects": [{"objectCode": "t"}, "not-a-record"],
            }
        )

        assert all(isinstance(step, str) for step in chain["reasoningSteps"])
        assert all(isinstance(item, dict) for item in chain["proposedObjects"])
        assert chain["reasoningSteps"] == ["一", "二"]
        assert chain["proposedObjects"] == [{"objectCode": "t"}]

    def test_the_endpoint_reports_which_fields_were_coerced(self, tmp_path: Path) -> None:
        platform_db, source = _scanned_source(tmp_path)
        result = generate_ontology_reasoning_chain(
            platform_db,
            source.id,
            _blueprint_client(
                {
                    "summary": "模型摘要",
                    "reasoningSteps": [{"step": 1, "detail": "读取元数据"}],
                    "proposedObjects": ["contract"],
                    "buildPlan": ["自动路径"],
                    "questionsForUser": ["哪些是主业务对象？"],
                },
                [],
            ),
        )

        assert result["usedRemoteModel"] is True
        assert result["unreadableModelFields"] == ["proposedObjects"]
        assert result["chain"]["reasoningSteps"] == ["读取元数据"]


def test_a_file_named_docx_that_is_not_one_is_refused_with_a_reason() -> None:
    """`docx` raises `PackageNotFoundError`, which is not a ValueError.

    Every caller's `except ValueError` therefore missed it, and uploading a file that was
    renamed to `.docx` -- or truncated in transit -- answered 500. The suffix check cannot
    catch this: it reads the name, and the name is exactly what is wrong.
    """
    with pytest.raises(ValueError) as refused:
        parse_rule_docx_bytes("rules.docx", b"this is not a Word document")

    message = str(refused.value)
    assert "Word" in message
    assert "PackageNotFound" not in message
    assert type(refused.value).__name__ == "ValueError"


def test_readiness_counts_mappings_once_per_mapping_not_once_per_object(tmp_path: Path) -> None:
    """A count over a join that fans out must count the mapping, not the join.

    The readiness summary joined `semantic_mapping` to `business_object` on `ontology_id`,
    so every mapping was matched once per object in its ontology. With four objects, a
    source with 32 pending mappings reported 128 -- and that figure is not merely displayed
    ("已确认 N 条映射，待审核 M 条"), it is the number the readiness panel shows. Coverage
    counts the same thing per object and sums, so two screens disagreed about how much was
    waiting for review. The count must be independent of how many objects the ontology has.
    """
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"
    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)

    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id)
    ontology_id = int(ontology["ontology"]["id"])

    with connect(platform_db) as conn:
        object_count = conn.execute(
            "select count(*) as n from business_object where ontology_id = ?", (ontology_id,)
        ).fetchone()["n"]
        pending = conn.execute(
            "select count(*) as n from semantic_mapping where ontology_id = ? and status = 'pending'",
            (ontology_id,),
        ).fetchone()["n"]

    assert object_count > 1, "这个用例的前提是本体里有多个对象，否则 JOIN 不会放大"
    assert pending > 0

    readiness = assess_data_source_readiness(platform_db, source.id)
    coverage = build_semantic_coverage(platform_db, source.id)

    assert readiness["summary"]["pendingMappings"] == pending
    # And the two surfaces that answer the same question agree with each other.
    assert readiness["summary"]["pendingMappings"] == coverage["summary"]["pendingMappings"]

    # The growth is by object count, so the assertion above is only meaningful if the two
    # numbers would differ under the old query -- which they would by exactly this factor.
    assert readiness["summary"]["pendingMappings"] != pending * object_count


# -- Re-scanning a source is not allowed to destroy an approved model --


def _source_with_draft_ontology(tmp_path: Path) -> tuple[Path, int, int]:
    """A scanned source plus the draft ontology derived from it."""
    platform_db = tmp_path / "platform.sqlite3"
    legacy_db = tmp_path / "legacy_contracts.sqlite3"
    initialize_platform_db(platform_db)
    create_contract_sample_db(legacy_db)
    source = register_data_source(platform_db, "合同管理样例系统", "sqlite", str(legacy_db), domain="合同管理")
    scan_data_source(platform_db, source.id)
    ontology = generate_ontology_draft(platform_db, source.id)
    return platform_db, source.id, ontology["id"]


def _set_ontology_status(platform_db: Path, ontology_id: int, status: str) -> None:
    with connect(platform_db) as conn:
        conn.execute("update ontology set status = ? where id = ?", (status, ontology_id))


def test_rescanning_keeps_a_published_ontology_intact(tmp_path: Path) -> None:
    """A published ontology is the artifact the governance story rests on.

    Re-scanning used to delete every object, attribute, rule and mapping behind it and
    leave the ontology row standing, still marked published -- so the damage was invisible
    in the one place a reader looks. The bindings are foreign keys, which is why the fix is
    to update the scanned rows in place rather than rebuild them: rebuilding orphans every
    object anchored to them.
    """
    platform_db, source_id, ontology_id = _source_with_draft_ontology(tmp_path)
    _set_ontology_status(platform_db, ontology_id, "published")

    def fingerprint() -> tuple:
        with connect(platform_db) as conn:
            return (
                conn.execute("select count(*) from business_object").fetchone()[0],
                conn.execute("select count(*) from business_attribute").fetchone()[0],
                conn.execute("select count(*) from business_rule").fetchone()[0],
                conn.execute("select count(*) from semantic_mapping").fetchone()[0],
                conn.execute("select id, status from ontology").fetchone()[:],
                # The identities matter as much as the counts: an object whose
                # `source_table_id` now points at nothing is a broken binding that a count
                # would not show.
                conn.execute("select source_table_id from business_object order by id").fetchall()[:],
            )

    before = fingerprint()
    assert before[0] > 0, "前提：草案里应当有业务对象"

    scan_data_source(platform_db, source_id)

    assert fingerprint() == before, "重新扫描不得改动已发布本体的任何绑定"


def test_a_scan_still_refuses_to_forget_a_column_a_model_maps_to(tmp_path: Path) -> None:
    """Forgetting a mapped column is the one thing a re-scan must not do silently.

    Refused at the point of damage rather than up front, which is what lets the ordinary
    re-scan -- same schema, or a new column -- proceed on a source that has a published
    ontology at all.
    """
    platform_db, source_id, ontology_id = _source_with_draft_ontology(tmp_path)
    _set_ontology_status(platform_db, ontology_id, "published")

    # Drop a column the published model maps, and remove it from the metadata the way a
    # re-scan would after the source really lost it.
    legacy_db = tmp_path / "legacy_contracts.sqlite3"
    with sqlite3.connect(legacy_db) as raw:
        raw.execute("alter table contract drop column title")

    with pytest.raises(ValueError) as refusal:
        scan_data_source(platform_db, source_id)

    message = str(refusal.value)
    assert "title" in message, message
    assert "业务属性" in message, message


def test_rescanning_discards_a_draft_and_what_later_modules_attached_to_it(tmp_path: Path) -> None:
    """The derived draft is regenerated, so everything hanging off it goes with it.

    Re-scanning a draft that had an event type attached failed with
    `FOREIGN KEY constraint failed`: the cleanup knew about rules, mappings, relations,
    attributes and objects, and nothing about the tables later feature modules added.
    """
    platform_db, source_id, ontology_id = _source_with_draft_ontology(tmp_path)
    # The event and temporal schemas are optional feature bundles: a deployment that never
    # uses them never creates their tables, which is why the cleanup has to ask the catalog
    # rather than assume they are there.
    with connect(platform_db) as conn:
        init_event_schema(conn)
        init_temporal_schema(conn)

    declare_event_type(
        platform_db,
        ontology_id,
        EventType(code="signed", name="签订", object_code="contract"),
        actor="test",
    )
    record_event(platform_db, ontology_id, "contract", "1", "signed", actor="test")
    record_attribute_version(platform_db, ontology_id, "contract", "1", "amount", 1200000, actor="test")

    scan_data_source(platform_db, source_id)

    with connect(platform_db) as conn:
        for table in ("ontology", "business_object", "business_event_type", "business_event", "attribute_version"):
            remaining = conn.execute(f"select count(*) from {table}").fetchone()[0]
            assert remaining == 0, f"{table} 在重新扫描后仍留有 {remaining} 行"


def test_rescanning_a_draft_from_an_unrelated_source_leaves_it_alone(tmp_path: Path) -> None:
    """The cleanup is scoped to the source being scanned.

    A cleanup query that is too broad deletes other people's work, and the symptom is an
    ontology that quietly loses its objects the next time somebody else's source is
    re-scanned.
    """
    platform_db = tmp_path / "platform.sqlite3"
    initialize_platform_db(platform_db)
    create_contract_sample_db(tmp_path / "legacy_contracts.sqlite3")

    # Two sources over the same database file, so both have objects on the same tables.
    first = register_data_source(
        platform_db, "源一", "sqlite", str(tmp_path / "legacy_contracts.sqlite3"), domain="合同"
    )
    second = register_data_source(
        platform_db, "源二", "sqlite", str(tmp_path / "legacy_contracts.sqlite3"), domain="合同"
    )
    scan_data_source(platform_db, first.id)
    scan_data_source(platform_db, second.id)
    ontology = generate_ontology_draft(platform_db, first.id)
    _set_ontology_status(platform_db, ontology["id"], "published")

    scan_data_source(platform_db, second.id)

    with connect(platform_db) as conn:
        kept = conn.execute("select count(*) from business_object where ontology_id = ?", (ontology["id"],)).fetchone()[
            0
        ]
    assert kept > 0, "扫描另一个数据源不得动到这个本体"
